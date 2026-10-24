"""Offline, deterministic integration smoke for the prepared development image."""

import json
import shutil
import subprocess
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pypdf import PdfReader
from reportlab.pdfgen.canvas import Canvas

plt.switch_backend("Agg")


def command(*argv):
    return subprocess.run(argv, check=True, capture_output=True, text=True).stdout.strip()


out = Path("output")
out.mkdir()
app = FastAPI()


@app.get("/health")
def health():
    return {"ready": True}


assert TestClient(app).get("/health").json() == {"ready": True}
frame = pd.DataFrame({"value": [1, 2, 3]})
assert frame["value"].mean() == np.mean([1, 2, 3]) == 2
frame.to_csv(out / "data.csv", index=False)
plt.plot(frame["value"])
plt.savefig(out / "plot.png")
plt.close()
book = Workbook()
book.active["A1"] = "offline workbook"
book.save(out / "book.xlsx")
assert load_workbook(out / "book.xlsx").active["A1"].value == "offline workbook"
doc = Document()
doc.add_paragraph("offline document")
doc.save(out / "document.docx")
assert Document(out / "document.docx").paragraphs[0].text == "offline document"
slides = Presentation()
slides.slides.add_slide(slides.slide_layouts[6])
slides.save(out / "slides.pptx")
assert len(Presentation(out / "slides.pptx").slides) == 1
pdf = Canvas(str(out / "document.pdf"))
pdf.drawString(30, 700, "offline PDF")
pdf.save()
assert "offline PDF" in PdfReader(out / "document.pdf").pages[0].extract_text()
assert "offline PDF" in command("pdftotext", str(out / "document.pdf"), "-")
image = Image.new("RGB", (700, 100), "white")
ImageDraw.Draw(image).text(
    (15, 15), "SANDBOX READY", fill="black", font=ImageFont.load_default(size=40)
)
image.save(out / "ocr.png")
assert "SANDBOX" in command("tesseract", str(out / "ocr.png"), "stdout", "--psm", "7")
command(
    "ffmpeg",
    "-v",
    "error",
    "-f",
    "lavfi",
    "-i",
    "sine=frequency=440:duration=0.1",
    str(out / "tone.wav"),
)
(out / "graph.dot").write_text("digraph { input -> output }")
command("dot", "-Tpng", str(out / "graph.dot"), "-o", str(out / "graph.png"))
(out / "main.c").write_text('#include <stdio.h>\nint main(void){puts("compiled");return 0;}\n')
command("gcc", str(out / "main.c"), "-o", str(out / "compiled"))
assert command(str(out / "compiled")) == "compiled"
assert command("sqlite3", ":memory:", "SELECT 6*7;") == "42"
(out / "main.ts").write_text("const value: number = 42; console.log(value)")
assert command("tsx", str(out / "main.ts")) == "42"
assert command("node", "-e", "console.log(require.resolve('typescript'))")
for name in ("python.freeze.txt", "node.freeze.json", "system.freeze.txt"):
    shutil.copyfile("/opt/validation-" + name, out / name)
(out / "smoke.json").write_text(
    json.dumps(
        {
            "passed": [
                "fastapi",
                "data",
                "plots",
                "xlsx",
                "docx",
                "pptx",
                "pdf",
                "ocr",
                "ffmpeg",
                "graphviz",
                "gcc",
                "sqlite",
                "typescript",
            ]
        },
        indent=2,
    )
)
print("toolbox smoke passed")
